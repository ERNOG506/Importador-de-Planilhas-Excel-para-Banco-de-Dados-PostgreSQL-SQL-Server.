from pathlib import Path

from openpyxl import load_workbook

from app.core.config import REQUIRED_COLUMNS


class ExcelReader:
    def read(self, file_path: str) -> list[dict]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Arquivo nao encontrado: {file_path}")
        if path.suffix.lower() != ".xlsx":
            raise ValueError("O arquivo deve estar no formato .xlsx")

        workbook = load_workbook(path, data_only=True)
        worksheet = workbook.active

        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("A planilha esta vazia")

        headers = [str(value).strip().lower() if value is not None else "" for value in rows[0]]
        missing_columns = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing_columns:
            missing = ", ".join(missing_columns)
            raise ValueError(f"Colunas obrigatorias ausentes: {missing}")

        records = []
        for row_number, row in enumerate(rows[1:], start=2):
            if all(value is None for value in row):
                continue
            record = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
            record["linha"] = row_number
            records.append(record)

        return records
